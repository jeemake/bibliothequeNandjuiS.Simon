# -*- coding: utf-8 -*-
import sys, os, html
sys.path.insert(0, "/workspace/decimo_final/build")
import layout as L
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.colors import HexColor
from pypdf import PdfReader, PdfWriter, Transformation, PageObject
FD = "/workspace/decimo_final/polices/"
FONTS = {"serif": "archivo-latin-600-normal", "jost3": "archivo-latin-400-normal", "jost4": "archivo-latin-400-normal", "jost5": "archivo-latin-500-normal", "manrope5": "manrope-latin-500-normal"}
for k, f in FONTS.items(): pdfmetrics.registerFont(TTFont(k, FD + f + ".ttf"))
PPT_FONT = {"serif": "Archivo", "jost3": "Archivo", "jost4": "Archivo", "jost5": "Archivo", "manrope5": "Manrope"}

def pstyle(st):
    return ParagraphStyle("s", fontName=st["font"], fontSize=st["size"], leading=st["lead"], textColor=HexColor(st["color"]),
                          alignment=TA_RIGHT if st.get("align") == "right" else TA_LEFT)
def ptext(t, st):
    t = t.upper() if st.get("caps") else t
    return html.escape(t).replace("\n", "<br/>")
def measure(t, st, w):
    p = Paragraph(ptext(t, st), pstyle(st)); _, h = p.wrap(w * mm, 5000); return h / mm
L._measure = measure
P = L.pages()

def draw_pdf(els, path):
    c = canvas.Canvas(path, pagesize=(L.W * mm, L.H * mm))
    c.setFillColor(HexColor(L.CHAUX)); c.rect(0, 0, L.W * mm, L.H * mm, stroke=0, fill=1)
    for el in els:
        k = el[0]
        if k == "rect":
            _, x, y, w, h, f = el; c.setFillColor(HexColor(f)); c.rect(x * mm, (L.H - y - h) * mm, w * mm, h * mm, stroke=0, fill=1)
        elif k == "img":
            _, p, x, y, w, h = el; c.drawImage(p, x * mm, (L.H - y - h) * mm, w * mm, h * mm, mask="auto")
        elif k == "line":
            _, x1, y1, x2, y2, col, pt = el; c.setStrokeColor(HexColor(col)); c.setLineWidth(pt); c.line(x1 * mm, (L.H - y1) * mm, x2 * mm, (L.H - y2) * mm)
        elif k == "text":
            _, x, y, w, t, st = el; p = Paragraph(ptext(t, st), pstyle(st)); _, h = p.wrap(w * mm, 5000); p.drawOn(c, x * mm, (L.H - y) * mm - h)
    c.showPage(); c.save()

def build_pdf(out):
    wr = PdfWriter(); tmp = []
    for i, pg in enumerate(P):
        if pg[0] == "page":
            f = f"{L.OUT}/_p{i:02d}.pdf"; draw_pdf(pg[1], f); r = PdfReader(f); wr.add_page(r.pages[0])
        else:
            r = PdfReader(pg[1]); src = r.pages[0]
            sw, sh = float(src.mediabox.width), float(src.mediabox.height)
            tw, th = L.W * mm, L.H * mm
            if abs(sw - tw) > 2 or abs(sh - th) > 2:
                s = min(tw / sw, th / sh); blank = PageObject.create_blank_page(width=tw, height=th)
                blank.merge_transformed_page(src, Transformation().scale(s).translate((tw - sw * s) / 2, (th - sh * s) / 2)); wr.add_page(blank)
            else:
                wr.add_page(src)
    wr.add_metadata({"/Title": "Bibliothèque Decimo · Esquisse V5 · dossier de présentation · indice F", "/Author": "K&Démi"})
    wr.compress_identical_objects(); wr.write(out)

def build_pptx(out):
    from pptx import Presentation
    from pptx.util import Mm, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    pr = Presentation(); pr.slide_width = Mm(L.W); pr.slide_height = Mm(L.H)
    blank = pr.slide_layouts[6]
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))
    for pg in P:
        s = pr.slides.add_slide(blank)
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(L.CHAUX)
        if pg[0] == "pdf":
            s.shapes.add_picture(pg[2], 0, 0, Mm(L.W), Mm(L.H)); continue
        for el in pg[1]:
            k = el[0]
            if k == "rect":
                _, x, y, w, h, f = el; sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y), Mm(w), Mm(h))
                sh.fill.solid(); sh.fill.fore_color.rgb = rgb(f); sh.line.fill.background(); sh.shadow.inherit = False
            elif k == "img":
                _, p, x, y, w, h = el; s.shapes.add_picture(p, Mm(x), Mm(y), Mm(w), Mm(h))
            elif k == "line":
                _, x1, y1, x2, y2, col, pt = el; ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
                ln.line.color.rgb = rgb(col); ln.line.width = Pt(pt)
            elif k == "text":
                _, x, y, w, t, st = el; h = measure(t, st, w) + 1.5
                tb = s.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h)); tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.TOP
                tt = t.upper() if st.get("caps") else t
                for i, line in enumerate(tt.split("\n")):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = PP_ALIGN.RIGHT if st.get("align") == "right" else PP_ALIGN.LEFT
                    para.line_spacing = Pt(st["lead"])
                    r = para.add_run(); r.text = line; r.font.name = PPT_FONT[st["font"]]; r.font.size = Pt(st["size"]); r.font.color.rgb = rgb(st["color"]); r.font.bold = True if st["font"] == "serif" else None
    pr.save(out)

if __name__ == "__main__":
    build_pdf(L.OUT + "/Decimo_ESQ_V5_Dossier_A3_v6.pdf")
    if "--pptx" in sys.argv: build_pptx(L.OUT + "/Decimo_ESQ_V5_Dossier_A3_v6.pptx")
    print("pages", len(P))
